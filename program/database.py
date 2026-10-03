# ==========================================================
# DATABASE.PY
#
# Tugas:
# 1. Membuat file Excel
# 2. Menyimpan data absensi
# 3. Menyimpan foto absensi
# 4. Menghitung jumlah absensi
# ==========================================================

import os

import cv2

from openpyxl import (
    Workbook,
    load_workbook
)

from openpyxl.drawing.image import Image


# ==========================================================
# MENENTUKAN LOKASI FOLDER
# ==========================================================

# Mendapatkan folder Face_id
BASE_DIR = os.path.dirname(
    os.path.dirname(
        os.path.abspath(__file__)
    )
)


# Folder foto absensi
FOLDER_ABSEN = os.path.join(
    BASE_DIR,
    "foto_absen"
)


# File Excel
FILE_EXCEL = os.path.join(
    BASE_DIR,
    "absensi.xlsx"
)


# Membuat folder foto jika belum ada
os.makedirs(
    FOLDER_ABSEN,
    exist_ok=True
)


# ==========================================================
# MEMBUAT EXCEL
# ==========================================================

def buat_excel():

    # Jika Excel belum ada
    if not os.path.exists(FILE_EXCEL):

        # Membuat workbook baru
        workbook = Workbook()


        # Mengambil sheet aktif
        sheet = workbook.active


        # Memberi nama sheet
        sheet.title = "Absensi"


        # Membuat judul kolom
        sheet["A1"] = "Nama"
        sheet["B1"] = "Tanggal"
        sheet["C1"] = "Jam"
        sheet["D1"] = "Sesi"
        sheet["E1"] = "Foto"


        # Menyimpan Excel
        workbook.save(
            FILE_EXCEL
        )


        # Menutup workbook
        workbook.close()


# Membuat Excel ketika program dimulai
buat_excel()


# ==========================================================
# MENGHITUNG ABSENSI
# ==========================================================

def hitung_absensi(
    nama,
    tanggal,
    sesi
):

    # Membuka Excel
    workbook = load_workbook(
        FILE_EXCEL
    )


    # Membuka sheet Absensi
    sheet = workbook["Absensi"]


    # Penghitung absensi
    jumlah = 0


    # Membaca setiap baris
    for baris in sheet.iter_rows(
        min_row=2,
        values_only=True
    ):

        nama_excel = baris[0]
        tanggal_excel = baris[1]
        sesi_excel = baris[3]


        # Membandingkan data
        if (
            nama_excel == nama
            and str(tanggal_excel) == str(tanggal)
            and sesi_excel == sesi
        ):

            jumlah += 1


    # Menutup Excel
    workbook.close()


    return jumlah


# ==========================================================
# MENYIMPAN ABSENSI
# ==========================================================

def simpan_absensi(
    nama,
    tanggal,
    jam,
    sesi,
    nomor_absen,
    frame
):

    # ------------------------------------------------------
    # MEMBUAT NAMA FOTO
    # ------------------------------------------------------

    nama_bersih = nama.replace(
        " ",
        "_"
    )


    nama_file = (
        f"{nama_bersih}_"
        f"{tanggal}_"
        f"{sesi}_"
        f"{nomor_absen}.jpg"
    )


    # Lokasi foto
    lokasi_foto = os.path.join(
        FOLDER_ABSEN,
        nama_file
    )


    # ------------------------------------------------------
    # MENYIMPAN FOTO
    # ------------------------------------------------------

    cv2.imwrite(
        lokasi_foto,
        frame
    )


    # ------------------------------------------------------
    # MEMBUKA EXCEL
    # ------------------------------------------------------

    workbook = load_workbook(
        FILE_EXCEL
    )


    sheet = workbook["Absensi"]


    # Baris kosong berikutnya
    baris = sheet.max_row + 1


    # ------------------------------------------------------
    # MENYIMPAN DATA
    # ------------------------------------------------------

    sheet.cell(
        baris,
        1,
        nama
    )

    sheet.cell(
        baris,
        2,
        tanggal
    )

    sheet.cell(
        baris,
        3,
        jam
    )

    sheet.cell(
        baris,
        4,
        sesi
    )


    # ------------------------------------------------------
    # MEMASUKKAN FOTO KE EXCEL
    # ------------------------------------------------------

    try:

        gambar_excel = Image(
            lokasi_foto
        )


        # Mengecilkan foto
        gambar_excel.width = 120
        gambar_excel.height = 90


        # Meletakkan foto di kolom E
        sheet.add_image(
            gambar_excel,
            f"E{baris}"
        )


        # Membuat tinggi baris lebih besar
        sheet.row_dimensions[
            baris
        ].height = 70


    except Exception as error:

        print(
            "Foto tidak dapat dimasukkan "
            f"ke Excel: {error}"
        )


        # Jika gagal memasukkan foto,
        # simpan nama file sebagai gantinya
        sheet.cell(
            baris,
            5,
            nama_file
        )


    # ------------------------------------------------------
    # MENYIMPAN EXCEL
    # ------------------------------------------------------

    workbook.save(
        FILE_EXCEL
    )


    workbook.close()


    print(
        "Data berhasil disimpan ke Excel."
    )

    print(
        f"Foto tersimpan: {lokasi_foto}"
    )